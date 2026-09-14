from pathlib import Path
from types import SimpleNamespace

import fitz
import pytest
from docx import Document

from app.models.cv import CVSourceType
from app.services.cv_text_extraction_service import CVTextExtractionService
from app.services.document_text_extractor import (
    DOCXTextExtractor,
    DocumentExtractionError,
    DocumentFileType,
    DocumentStorageError,
    DocumentTextTooLargeError,
    EmptyDocumentTextError,
    EncryptedDocumentError,
    PDFTextExtractor,
    UnsupportedDocumentTypeError,
    UploadedDocumentTextExtractor,
    normalize_document_text,
)


PDF_TEXT = "Nguyen Van A\nBackend Developer\nPython FastAPI PostgreSQL"


def _write_pdf(path: Path, text: str = PDF_TEXT) -> Path:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(path)
    document.close()
    return path


def _write_vietnamese_pdf(path: Path) -> Path:
    document = fitz.open()
    page = document.new_page()
    page.insert_htmlbox(
        fitz.Rect(72, 72, 500, 300),
        "<p>Nguyễn Văn A<br>Đại học Công nghệ<br>Kỹ năng: Python</p>",
    )
    document.save(path)
    document.close()
    return path


def _write_docx(path: Path) -> Path:
    document = Document()
    document.add_paragraph("Nguyễn Văn A")
    document.add_paragraph("Kinh nghiệm: Backend Developer")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Kỹ năng"
    table.cell(0, 1).text = "Python, FastAPI, PostgreSQL"
    table.cell(1, 0).text = "Học vấn"
    table.cell(1, 1).text = "Đại học Công nghệ"
    document.save(path)
    return path


def test_pdf_extracts_semantic_text_and_page_count(tmp_path: Path) -> None:
    result = PDFTextExtractor().extract(_write_pdf(tmp_path / "candidate.pdf"))

    assert result.file_type == DocumentFileType.PDF
    assert result.page_count == 1
    assert result.character_count == len(result.text)
    for value in ("Nguyen Van A", "Backend Developer", "Python", "FastAPI", "PostgreSQL"):
        assert value in result.text


def test_pdf_preserves_vietnamese_unicode(tmp_path: Path) -> None:
    result = PDFTextExtractor().extract(_write_vietnamese_pdf(tmp_path / "unicode.pdf"))
    assert "Nguyễn Văn A" in result.text
    assert "Đại học Công nghệ" in result.text
    assert "Kỹ năng: Python" in result.text


def test_docx_preserves_vietnamese_paragraphs_and_table_text(tmp_path: Path) -> None:
    result = DOCXTextExtractor().extract(_write_docx(tmp_path / "candidate.docx"))

    assert result.file_type == DocumentFileType.DOCX
    assert result.page_count is None
    for value in (
        "Nguyễn Văn A",
        "Kinh nghiệm",
        "Backend Developer",
        "Kỹ năng | Python, FastAPI, PostgreSQL",
        "Học vấn | Đại học Công nghệ",
    ):
        assert value in result.text


def test_normalization_preserves_unicode_and_cleans_whitespace() -> None:
    raw = "\x00Nguyễn   Văn\tA  \r\n\r\n\r\n  Kỹ năng:   Python  \n"
    assert normalize_document_text(raw) == "Nguyễn Văn A\n\nKỹ năng: Python"


def test_empty_valid_pdf_reports_ocr_is_not_supported(tmp_path: Path) -> None:
    document = fitz.open()
    document.new_page()
    path = tmp_path / "scanned.pdf"
    document.save(path)
    document.close()

    with pytest.raises(EmptyDocumentTextError, match="OCR chưa được hỗ trợ"):
        PDFTextExtractor().extract(path)


def test_encrypted_pdf_is_rejected_safely(tmp_path: Path) -> None:
    source = fitz.open()
    source.new_page().insert_text((72, 72), PDF_TEXT)
    path = tmp_path / "protected.pdf"
    source.save(
        path,
        encryption=fitz.PDF_ENCRYPT_AES_256,
        owner_pw="owner-secret",
        user_pw="user-secret",
    )
    source.close()

    with pytest.raises(EncryptedDocumentError, match="mật khẩu"):
        PDFTextExtractor().extract(path)


@pytest.mark.parametrize(
    ("extractor", "filename"),
    [(PDFTextExtractor(), "broken.pdf"), (DOCXTextExtractor(), "broken.docx")],
)
def test_corrupt_documents_map_parser_errors(extractor, filename: str, tmp_path: Path) -> None:
    path = tmp_path / filename
    path.write_bytes(b"not a valid document")

    with pytest.raises(DocumentExtractionError, match="bị hỏng"):
        extractor.extract(path)


def test_dispatch_requires_matching_validated_mime_and_extension(tmp_path: Path) -> None:
    path = _write_pdf(tmp_path / "candidate.pdf")
    dispatcher = UploadedDocumentTextExtractor()

    with pytest.raises(UnsupportedDocumentTypeError, match="PDF hoặc DOCX"):
        dispatcher.extract(path, "text/plain", "candidate.pdf")
    with pytest.raises(UnsupportedDocumentTypeError, match="PDF hoặc DOCX"):
        dispatcher.extract(path, "application/pdf", "candidate.docx")


def test_extracted_text_limit_fails_without_truncating(tmp_path: Path) -> None:
    path = _write_pdf(tmp_path / "large.pdf", "\n".join(["A" * 50] * 4))
    with pytest.raises(DocumentTextTooLargeError, match="giới hạn"):
        PDFTextExtractor(max_characters=100).extract(path)


def test_uploaded_cv_service_reuses_storage_and_returns_text(tmp_path: Path) -> None:
    path = _write_docx(tmp_path / "stored.docx")
    storage = SimpleNamespace(get_file_path=lambda storage_key: path)
    cv = SimpleNamespace(
        source_type=CVSourceType.UPLOADED,
        storage_key="private/internal/key.docx",
        mime_type=UploadedDocumentTextExtractor.DOCX_MIME,
        original_filename="candidate.docx",
    )

    result = CVTextExtractionService(storage=storage).extract_uploaded_cv(cv)
    assert "Nguyễn Văn A" in result.text


def test_builder_cv_is_never_sent_to_document_parser() -> None:
    storage = SimpleNamespace(get_file_path=lambda storage_key: pytest.fail("storage was read"))
    cv = SimpleNamespace(source_type=CVSourceType.BUILDER)

    with pytest.raises(UnsupportedDocumentTypeError, match="Builder"):
        CVTextExtractionService(storage=storage).extract_uploaded_cv(cv)


def test_storage_failure_is_mapped_without_leaking_path() -> None:
    def fail_read(storage_key):
        raise FileNotFoundError("C:/private/candidate-secret.pdf")

    storage = SimpleNamespace(get_file_path=fail_read)
    cv = SimpleNamespace(
        source_type=CVSourceType.UPLOADED,
        storage_key="private/key.pdf",
        mime_type="application/pdf",
        original_filename="candidate.pdf",
    )

    with pytest.raises(DocumentStorageError) as error:
        CVTextExtractionService(storage=storage).extract_uploaded_cv(cv)
    assert "candidate-secret" not in str(error.value)
